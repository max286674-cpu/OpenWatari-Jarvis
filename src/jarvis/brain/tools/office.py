"""Direct local Office/data tools.

Use deterministic file APIs for Word, Excel and SQLite instead of GUI automation whenever possible.
"""
from __future__ import annotations
import json
import sqlite3
from pathlib import Path
from jarvis.brain.tools.base import tool_error

def _path(value: str) -> Path:
    p = Path(value).expanduser()
    if not p.is_absolute(): p = Path.cwd() / p
    return p.resolve()

async def create_word(args: dict) -> str:
    try:
        from docx import Document
        from docx.shared import Cm, Pt
        path = _path(str(args.get("path") or "document.docx")); path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document(); sec = doc.sections[0]
        sec.top_margin=Cm(float(args.get("top_cm",2))); sec.bottom_margin=Cm(float(args.get("bottom_cm",2)))
        sec.left_margin=Cm(float(args.get("left_cm",3))); sec.right_margin=Cm(float(args.get("right_cm",1.5)))
        style=doc.styles["Normal"]; style.font.name=str(args.get("font","Times New Roman")); style.font.size=Pt(float(args.get("font_size",14)))
        for block in str(args.get("text") or "").split("\n\n"): doc.add_paragraph(block)
        doc.save(path); return f"Создан Word-документ: {path}"
    except Exception as e: return tool_error("create Word document", e)

async def edit_word(args: dict) -> str:
    try:
        from docx import Document
        path=_path(str(args.get("path") or ""))
        if not path.exists(): return f"Файл не найден: {path}"
        doc=Document(path); find=str(args.get("find") or ""); repl=str(args.get("replace") or ""); changed=0
        if find:
            for p in doc.paragraphs:
                if find in p.text:
                    for r in p.runs:
                        if find in r.text: r.text=r.text.replace(find,repl); changed+=1
        if args.get("append"): doc.add_paragraph(str(args["append"])); changed+=1
        doc.save(path); return f"Word-документ обновлён: {path}; изменений: {changed}"
    except Exception as e: return tool_error("edit Word document", e)

async def create_excel(args: dict) -> str:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
        path=_path(str(args.get("path") or "table.xlsx")); path.parent.mkdir(parents=True,exist_ok=True)
        wb=Workbook(); ws=wb.active; ws.title=str(args.get("sheet") or "Лист1")
        rows=args.get("rows") or []; rows=json.loads(rows) if isinstance(rows,str) else rows
        for i,row in enumerate(rows,1):
            for j,value in enumerate(row,1): ws.cell(i,j,value)
        if rows:
            for cell in ws[1]: cell.font=Font(bold=True); cell.alignment=Alignment(horizontal="center")
            ws.freeze_panes="A2"
        wb.save(path); return f"Создан Excel-файл: {path}"
    except Exception as e: return tool_error("create Excel file", e)

async def edit_excel(args: dict) -> str:
    try:
        from openpyxl import load_workbook
        path=_path(str(args.get("path") or ""))
        if not path.exists(): return f"Файл не найден: {path}"
        wb=load_workbook(path); ws=wb[str(args.get("sheet") or wb.sheetnames[0])]
        if args.get("cell"): ws[str(args["cell"])]=args.get("value")
        for item in args.get("cells") or []: ws[str(item["cell"])]=item.get("value")
        wb.save(path); return f"Excel-файл обновлён: {path}"
    except Exception as e: return tool_error("edit Excel file", e)

async def create_database(args: dict) -> str:
    try:
        path=_path(str(args.get("path") or "database.sqlite")); path.parent.mkdir(parents=True,exist_ok=True)
        schema=args.get("schema") or []; schema=json.loads(schema) if isinstance(schema,str) else schema
        with sqlite3.connect(path) as con:
            for table in schema:
                name=str(table["name"]); cols=", ".join(str(c) for c in table["columns"])
                con.execute(f'CREATE TABLE IF NOT EXISTS "{name}" ({cols})')
        return f"База данных создана: {path}"
    except Exception as e: return tool_error("create database", e)

SCHEMAS=[
 {"type":"function","function":{"name":"create_word","description":"Создать DOCX с текстом и базовым оформлением.","parameters":{"type":"object","properties":{"path":{"type":"string"},"text":{"type":"string"},"font":{"type":"string"},"font_size":{"type":"number"},"top_cm":{"type":"number"},"bottom_cm":{"type":"number"},"left_cm":{"type":"number"},"right_cm":{"type":"number"}},"required":["text"]}}},
 {"type":"function","function":{"name":"edit_word","description":"Редактировать DOCX: заменить текст или добавить абзац.","parameters":{"type":"object","properties":{"path":{"type":"string"},"find":{"type":"string"},"replace":{"type":"string"},"append":{"type":"string"}},"required":["path"]}}},
 {"type":"function","function":{"name":"create_excel","description":"Создать XLSX из двумерного массива строк.","parameters":{"type":"object","properties":{"path":{"type":"string"},"sheet":{"type":"string"},"rows":{"type":"array","items":{"type":"array"}}},"required":["rows"]}}},
 {"type":"function","function":{"name":"edit_excel","description":"Редактировать XLSX, записывая одну или несколько ячеек.","parameters":{"type":"object","properties":{"path":{"type":"string"},"sheet":{"type":"string"},"cell":{"type":"string"},"value":{},"cells":{"type":"array","items":{"type":"object"}}},"required":["path"]}}},
 {"type":"function","function":{"name":"create_database","description":"Создать SQLite-базу и таблицы по схеме.","parameters":{"type":"object","properties":{"path":{"type":"string"},"schema":{"type":"array","items":{"type":"object"}}},"required":["schema"]}}},
]
HANDLERS={"create_word":create_word,"edit_word":edit_word,"create_excel":create_excel,"edit_excel":edit_excel,"create_database":create_database}
